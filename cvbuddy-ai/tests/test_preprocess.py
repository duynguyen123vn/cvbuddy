from app.pipeline.preprocess import (
    compute_local_metrics,
    _group_pdf_blocks_by_column,
    extract_docx_document,
    extract_pdf_document,
    mask_pii,
    normalize_section_lines,
    normalize_text,
    split_into_sections,
)
from app.schemas.cv import CvSections
from app.pipeline.ai_extract import _remove_section_heading_items
from app.pipeline.extract import (
    TargetLanguage,
    reconcile_ai_sections,
    select_pdf_language_pages,
)


def test_mask_pii() -> None:
    text = "Contact: user@mail.com or 0912345678"
    masked = mask_pii(text)
    assert "[EMAIL]" in masked
    assert "[PHONE]" in masked


def test_group_pdf_blocks_keeps_two_columns_contiguous() -> None:
    blocks = [
        (24.0, 40.0, 180.0, 60.0, "Left heading"),
        (24.0, 100.0, 180.0, 120.0, "Left detail"),
        (320.0, 50.0, 570.0, 70.0, "Right heading"),
        (320.0, 110.0, 570.0, 130.0, "Right detail"),
    ]

    groups, column_count = _group_pdf_blocks_by_column(blocks, page_width=600.0)

    assert column_count == 2
    assert [block[4] for block in groups[0]] == ["Left heading", "Left detail"]
    assert [block[4] for block in groups[1]] == ["Right heading", "Right detail"]


def test_pymupdf_extracts_two_columns_in_column_order() -> None:
    import fitz

    document = fitz.open()
    page = document.new_page(width=600, height=800)
    page.insert_text((30, 60), "Left heading")
    page.insert_text((30, 120), "Left detail one")
    page.insert_text((30, 180), "Left detail two")
    page.insert_text((330, 70), "Right heading")
    page.insert_text((330, 130), "Right detail one")
    page.insert_text((330, 190), "Right detail two")
    pdf_bytes = document.tobytes()
    document.close()

    extracted = extract_pdf_document(pdf_bytes)

    assert extracted.max_columns == 2
    assert "[Column 1]" in extracted.normalized_text
    assert "[Column 2]" in extracted.normalized_text
    assert extracted.normalized_text.index("Left detail two") < extracted.normalized_text.index("Right heading")


def test_split_sections() -> None:
    raw = """
Experience
- Led a student event with 50 attendees
Education
- University ABC
"""
    sections = split_into_sections(raw)
    assert len(sections.experience) == 1
    assert "50" in sections.experience[0]


def test_project_word_at_end_of_bullet_is_not_treated_as_heading() -> None:
    sections = split_into_sections("Experience\nBuilt an AI project\nSkills\nPython")

    assert sections.experience == ["Built an AI project"]
    assert sections.skills == ["Python"]


def test_ai_section_cleanup_removes_headings_from_model_output() -> None:
    cleaned = _remove_section_heading_items(
        CvSections(
            summary=["Career Objective", "AI student"],
            education=["Education", "FPT University"],
            skills=["Skills", "Python"],
        )
    )

    assert cleaned.summary == ["AI student"]
    assert cleaned.education == ["FPT University"]
    assert cleaned.skills == ["Python"]


def test_normalize_repairs_common_pdf_mojibake() -> None:
    text = normalize_text("ThÃ nh viÃªn phÃ¡t triá»ƒn mÃ´ hÃ¬nh AI")
    assert "Thành viên" in text
    assert "mô hình" in text


def test_normalize_keeps_valid_vietnamese_with_pdf_icon_noise() -> None:
    text = normalize_text(" Nguyễn Hoàng Duy\nThành viên phát triển mô hình AI")
    assert "Nguyễn Hoàng Duy" in text
    assert "Thành viên phát triển mô hình AI" in text
    assert "" not in text


def test_split_sections_common_cv_headings() -> None:
    raw = """
Career Objective
Become an AI Engineer
Technical Skills
Python
Certificates
Machine Learning Specialization
"""
    sections = split_into_sections(raw)
    assert sections.summary == ["Become an AI Engineer"]
    assert sections.skills == ["Python"]
    assert sections.certifications == ["Machine Learning Specialization"]


def test_split_sections_vietnamese_cv_headings() -> None:
    raw = """
Mục tiêu nghề nghiệp
Trở thành kỹ sư AI
Học vấn
Đại học FPT
Kỹ năng
Python
Kinh nghiệm
Tăng độ chính xác mô hình lên 20%
"""
    sections = split_into_sections(raw)
    assert sections.summary == ["Trở thành kỹ sư AI"]
    assert sections.education == ["Đại học FPT"]
    assert sections.skills == ["Python"]
    assert sections.experience == ["Tăng độ chính xác mô hình lên 20%"]


def test_split_sections_handles_glued_two_column_headings() -> None:
    raw = """
Cai Rang Ward, Can Tho City Education
FPT University
Deep Learning Experience
Team Member | Face Recognition Attendance System
"""
    sections = split_into_sections(raw)
    assert "FPT University" in sections.education
    assert "Team Member | Face Recognition Attendance System" in sections.experience


def test_split_sections_handles_inline_heading_content() -> None:
    raw = """
Experience: Developed an automatic attendance system using Python and OpenCV.
Education FPT University
Skills Python, OpenCV, Machine Learning
"""
    sections = split_into_sections(raw)
    assert sections.experience == [
        "Developed an automatic attendance system using Python and OpenCV."
    ]
    assert sections.education == ["FPT University"]
    assert sections.skills == ["Python, OpenCV, Machine Learning"]


def test_split_sections_handles_project_heading_variants() -> None:
    raw = """
Relevant Projects
Face Recognition Attendance System
- Developed real-time attendance tracking with OpenCV.
Capstone Project: AI Resume Screening Assistant
Skills
Python
"""
    sections = split_into_sections(raw)
    assert "Face Recognition Attendance System" in sections.projects
    assert "Developed real-time attendance tracking with OpenCV." in sections.projects
    assert "AI Resume Screening Assistant" in sections.projects
    assert sections.skills == ["Python"]


def test_split_sections_handles_vietnamese_project_headings() -> None:
    raw = """
Dự án
Hệ thống điểm danh bằng nhận diện khuôn mặt
Đồ án: Phân loại CV bằng AI
"""
    sections = split_into_sections(raw)
    assert "Hệ thống điểm danh bằng nhận diện khuôn mặt" in sections.projects
    assert "Phân loại CV bằng AI" in sections.projects


def test_split_sections_handles_extended_heading_vocabulary_and_layout_markers() -> None:
    raw = """
[Page 1]
About Me
AI engineering student
Selected Work
Resume ranking platform
Tools and Technologies
Python, FastAPI, PostgreSQL
Training
Machine Learning Operations
"""
    sections = split_into_sections(raw)

    assert sections.summary == ["AI engineering student"]
    assert sections.projects == ["Resume ranking platform"]
    assert sections.skills == ["Python, FastAPI, PostgreSQL"]
    assert sections.certifications == ["Machine Learning Operations"]
    assert not any("[Page" in item for item in sections.other)


def test_docx_extraction_preserves_paragraph_and_table_order() -> None:
    from io import BytesIO

    from docx import Document

    doc = Document()
    doc.add_paragraph("Experience")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "2025"
    table.rows[0].cells[1].text = "AI Intern"
    doc.add_paragraph("Skills")
    doc.add_paragraph("Python")
    buffer = BytesIO()
    doc.save(buffer)

    result = extract_docx_document(buffer.getvalue())

    assert result.normalized_text.index("Experience") < result.normalized_text.index("2025")
    assert result.normalized_text.index("2025") < result.normalized_text.index("Skills")


def test_ai_reconciliation_restores_omitted_items_for_single_language_cv() -> None:
    ai_sections = CvSections(
        experience=["AI Intern"],
        skills=["Python"],
    )
    local_sections = CvSections(
        experience=["AI Intern", "Built a document parsing API"],
        skills=["Python"],
        projects=["Resume ranking platform"],
    )

    reconciled, restored = reconcile_ai_sections(
        ai_sections,
        local_sections,
        "Experience\nAI Intern\nBuilt a document parsing API\nProjects\nResume ranking platform\nSkills\nPython",
        TargetLanguage.AUTO,
    )

    assert restored == 2
    assert "Built a document parsing API" in reconciled.experience
    assert "Resume ranking platform" in reconciled.projects


def test_ai_reconciliation_does_not_merge_bilingual_versions_in_auto_mode() -> None:
    ai_sections = CvSections(summary=["AI engineering student"])
    local_sections = CvSections(
        summary=["AI engineering student", "Sinh viên kỹ thuật AI"],
        education=["FPT University", "Đại học FPT"],
    )
    text = """
Summary
AI engineering student
Education
FPT University
Mục tiêu nghề nghiệp
Sinh viên kỹ thuật AI
Học vấn
Đại học FPT
"""

    reconciled, restored = reconcile_ai_sections(
        ai_sections,
        local_sections,
        text,
        TargetLanguage.AUTO,
    )

    assert restored == 0
    assert reconciled == ai_sections


def test_select_pdf_language_pages_keeps_only_requested_cv_version() -> None:
    text = """
[Page 1]
Career Objective
Build AI products and develop machine learning systems.
Education
FPT University
[Page 2]
Mục tiêu nghề nghiệp
Phát triển sản phẩm AI và xây dựng hệ thống học máy.
Học vấn
Đại học FPT
"""

    vietnamese, vi_pages = select_pdf_language_pages(text, TargetLanguage.VI)
    english, en_pages = select_pdf_language_pages(text, TargetLanguage.EN)

    assert vi_pages == {2}
    assert "Mục tiêu nghề nghiệp" in vietnamese
    assert "Career Objective" not in vietnamese
    assert en_pages == {1}
    assert "Career Objective" in english
    assert "Mục tiêu nghề nghiệp" not in english


def test_normalize_section_lines_splits_multiple_glued_heading_shapes() -> None:
    lines = normalize_section_lines(
        [
            "Contact information Education FPT University",
            "Experience Developed a model",
            "Python Projects",
        ]
    )
    assert lines == [
        "Contact information",
        "Education",
        "FPT University",
        "Experience",
        "Developed a model",
        "Python",
        "Projects",
    ]


def test_local_metrics_quantified() -> None:
    sections = CvSections(experience=["Increased engagement by 20%"])
    m = compute_local_metrics(sections)
    assert m.has_quantified_bullets is True
    assert m.quantified_bullet_count == 1


def test_local_metrics_quantified_vietnamese() -> None:
    sections = CvSections(experience=["Tăng độ chính xác mô hình lên 20%"])
    m = compute_local_metrics(sections)
    assert m.has_quantified_bullets is True
    assert m.quantified_bullet_count == 1
