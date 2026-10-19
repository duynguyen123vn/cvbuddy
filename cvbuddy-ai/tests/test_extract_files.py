from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.main import app


@pytest.fixture
def client() -> TestClient:
    return TestClient(app)


def _docx_bytes(lines: list[str]) -> bytes:
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _table_docx_bytes(rows: list[tuple[str, str]]) -> bytes:
    doc = Document()
    table = doc.add_table(rows=len(rows), cols=2)
    for row_index, row_values in enumerate(rows):
        for col_index, value in enumerate(row_values):
            table.rows[row_index].cells[col_index].text = value

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


@pytest.mark.parametrize(
    ("filename", "lines", "expected"),
    [
        (
            "english-ai-cv.docx",
            [
                "Career Objective",
                "Become an AI Engineer",
                "Education",
                "FPT University",
                "Technical Skills",
                "Python",
                "Experience",
                "Improved model accuracy by 20%",
            ],
            {
                "summary": "Become an AI Engineer",
                "education": "FPT University",
                "skills": "Python",
                "experience": "Improved model accuracy by 20%",
                "quantified_bullet_count": 1,
            },
        ),
        (
            "vietnamese-ai-cv.docx",
            [
                "Mục tiêu nghề nghiệp",
                "Trở thành kỹ sư AI",
                "Học vấn",
                "Đại học FPT",
                "Kỹ năng",
                "Python",
                "Kinh nghiệm",
                "Tăng độ chính xác mô hình lên 20%",
            ],
            {
                "summary": "Trở thành kỹ sư AI",
                "education": "Đại học FPT",
                "skills": "Python",
                "experience": "Tăng độ chính xác mô hình lên 20%",
                "quantified_bullet_count": 1,
            },
        ),
        (
            "two-column-like-cv.docx",
            [
                "Cai Rang Ward, Can Tho City Education",
                "FPT University",
                "Deep Learning Experience",
                "Team Member | Face Recognition Attendance System",
                "Technical Skills",
                "Python, OpenCV",
            ],
            {
                "education": "FPT University",
                "skills": "Python, OpenCV",
                "experience": "Team Member | Face Recognition Attendance System",
                "quantified_bullet_count": 0,
            },
        ),
        (
            "inline-heading-cv.docx",
            [
                "Experience: Developed an automatic attendance system using Python and OpenCV.",
                "Education FPT University",
                "Skills Python, OpenCV, Machine Learning",
            ],
            {
                "education": "FPT University",
                "skills": "Python, OpenCV, Machine Learning",
                "experience": "Developed an automatic attendance system using Python and OpenCV.",
                "quantified_bullet_count": 0,
            },
        ),
        (
            "project-cv.docx",
            [
                "Relevant Projects",
                "Face Recognition Attendance System",
                "Developed real-time attendance tracking with OpenCV for 20 users",
                "Skills",
                "Python, OpenCV",
            ],
            {
                "skills": "Python, OpenCV",
                "projects": "Face Recognition Attendance System",
                "quantified_bullet_count": 1,
            },
        ),
    ],
)
def test_extract_docx_multiple_cv_shapes(
    client: TestClient,
    filename: str,
    lines: list[str],
    expected: dict[str, str | int],
) -> None:
    r = client.post(
        "/v1/cv/extract",
        files={
            "file": (
                filename,
                _docx_bytes(lines),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert r.status_code == 200
    data = r.json()
    sections = data["sections"]
    metrics = data["local_metrics"]
    assert {
        "bullet_count",
        "quantified_bullet_count",
        "has_quantified_bullets",
        "avg_bullet_length",
        "section_count",
        "detected_languages",
    } <= metrics.keys()

    for section_name in ("summary", "education", "skills", "experience", "projects"):
        if section_name in expected:
            assert expected[section_name] in sections[section_name]

    assert metrics["quantified_bullet_count"] == expected["quantified_bullet_count"]
    assert data["normalized_text_length"] > 0
    assert data["coverage"]["coverage_score"] == 1.0
    assert data["coverage"]["source_block_count"] > 0
    assert len(data["structured_document"]["raw_blocks"]) > 0


def test_extract_docx_reads_table_based_cv_template(client: TestClient) -> None:
    r = client.post(
        "/v1/cv/extract?mode=local&debug=true",
        files={
            "file": (
                "table-template-cv.docx",
                _table_docx_bytes(
                    [
                        ("Education", "FPT University"),
                        ("Technical Skills", "Python, OpenCV"),
                        ("Experience", "Developed attendance system with 20 users"),
                    ]
                ),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert r.status_code == 200
    data = r.json()
    assert "FPT University" in data["sections"]["education"]
    assert "Python, OpenCV" in data["sections"]["skills"]
    assert "Developed attendance system with 20 users" in data["sections"]["experience"]
    assert data["local_metrics"]["quantified_bullet_count"] == 1
    assert data["meta"]["document"]["engine"] == "python-docx"
    assert data["meta"]["document"]["text_block_count"] >= 6
    assert data["meta"]["document"]["needs_ocr"] is False
    assert any(
        block["source_type"] == "table_row"
        for block in data["structured_document"]["raw_blocks"]
    )


def test_lossless_extract_retains_unclassified_blocks(client: TestClient) -> None:
    response = client.post(
        "/v1/cv/extract?mode=local",
        files={
            "file": (
                "extra-content.docx",
                _docx_bytes(
                    [
                        "Experience",
                        "Built an attendance system",
                        "A standalone note the section parser cannot classify",
                    ]
                ),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["coverage"]["retained_block_count"] == payload["coverage"]["source_block_count"]
    assert payload["coverage"]["coverage_score"] == 1.0


def test_extract_debug_includes_normalized_text_preview(client: TestClient) -> None:
    r = client.post(
        "/v1/cv/extract?mode=local&debug=true",
        files={
            "file": (
                "debug-cv.docx",
                _docx_bytes(["Experience", "Built a project"]),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert r.status_code == 200
    assert "normalized_text_preview" in r.json()
    assert "Built a project" in r.json()["normalized_text_preview"]


def test_extract_response_includes_hybrid_meta(client: TestClient) -> None:
    r = client.post(
        "/v1/cv/extract?mode=hybrid",
        files={
            "file": (
                "hybrid-cv.docx",
                _docx_bytes(["Experience", "Built a project", "Skills", "Python"]),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert r.status_code == 200
    meta = r.json()["meta"]
    assert meta["extract_mode"] == "hybrid"
    assert meta["parser_source"] in {"local", "ai", "local_fallback"}
    assert isinstance(meta["local_quality_score"], float)
    assert meta["analysis_language"] in {"vi", "en", "both"}


def test_extract_ai_mode_uses_ai_parser(
    client: TestClient,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.pipeline import extract as extract_pipeline
    from app.schemas.cv import CvSections

    async def fake_ai_extract(
        _normalized_text: str,
        *,
        target_language: str = "auto",
        llm_model: str | None = None,
    ) -> CvSections:
        return CvSections(
            experience=["AI parsed experience"],
            education=["AI parsed education"],
            skills=["AI parsed skills"],
        )

    monkeypatch.setattr(extract_pipeline, "run_ai_extract_sections", fake_ai_extract)

    r = client.post(
        "/v1/cv/extract?mode=ai",
        files={
            "file": (
                "ai-cv.docx",
                _docx_bytes(["Unstructured text that local parser may not classify well."]),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert r.status_code == 200
    data = r.json()
    assert data["meta"]["parser_source"] == "ai"
    assert data["sections"]["experience"] == ["AI parsed experience"]


def test_extract_ai_mode_passes_target_language(
    client: TestClient,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.pipeline import extract as extract_pipeline
    from app.schemas.cv import CvSections

    seen: dict[str, str] = {}

    async def fake_ai_extract(
        _normalized_text: str,
        *,
        target_language: str = "auto",
        llm_model: str | None = None,
    ) -> CvSections:
        seen["target_language"] = target_language
        return CvSections(summary=["English version only"])

    monkeypatch.setattr(extract_pipeline, "run_ai_extract_sections", fake_ai_extract)

    r = client.post(
        "/v1/cv/extract?mode=ai&target_language=en",
        files={
            "file": (
                "bilingual-cv.docx",
                _docx_bytes(["Career Objective", "English text", "Mục tiêu nghề nghiệp", "Tiếng Việt"]),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert r.status_code == 200
    assert seen["target_language"] == "en"
    assert r.json()["meta"]["target_language"] == "en"
    assert r.json()["sections"]["summary"] == ["English version only"]


def test_extract_ai_mode_accepts_target_language_from_multipart_form(
    client: TestClient,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.pipeline import extract as extract_pipeline
    from app.schemas.cv import CvSections

    seen: dict[str, str] = {}

    async def fake_ai_extract(
        _normalized_text: str,
        *,
        target_language: str = "auto",
        llm_model: str | None = None,
    ) -> CvSections:
        seen["target_language"] = target_language
        return CvSections(summary=["Synthetic summary"])

    monkeypatch.setattr(extract_pipeline, "run_ai_extract_sections", fake_ai_extract)
    response = client.post(
        "/v1/cv/extract?mode=ai",
        data={"target_language": "both"},
        files={"file": ("synthetic.docx", _docx_bytes(["Synthetic CV text"]), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )

    assert response.status_code == 200
    assert seen["target_language"] == "both"
    assert response.json()["meta"]["target_language"] == "both"


def test_empty_extraction_is_structured_error(
    client: TestClient,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from types import SimpleNamespace

    from app.api.v1 import cv as cv_api
    from app.schemas.cv import CvSections

    monkeypatch.setattr(
        cv_api,
        "extract_cv_sections",
        lambda *_args, **_kwargs: _async_value(
            SimpleNamespace(normalized_text="", sections=CvSections())
        ),
    )
    response = client.post(
        "/v1/cv/extract?mode=local",
        files={"file": ("empty.pdf", b"synthetic", "application/pdf")},
    )

    assert response.status_code == 422
    assert response.headers["content-type"].startswith("application/json")
    assert response.json()["error"]["code"] == "validation_error"


def test_extract_and_analyze_rejects_empty_sections_before_analysis(
    client: TestClient,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from types import SimpleNamespace

    from app.api.v1 import cv as cv_api
    from app.pipeline import analyze as analyze_pipeline
    from app.schemas.cv import CvSections

    async def empty_extraction(*_args, **_kwargs):
        return SimpleNamespace(normalized_text="Synthetic PDF text", sections=CvSections())

    def fail_if_analyzed(*_args, **_kwargs):
        raise AssertionError("analysis must not run without extracted sections")

    monkeypatch.setattr(cv_api, "extract_cv_sections", empty_extraction)
    monkeypatch.setattr(analyze_pipeline, "run_cv_analyze", fail_if_analyzed)
    response = client.post(
        "/v1/cv/extract-and-analyze?industry_slug=ai_software",
        files={"file": ("synthetic.pdf", b"synthetic", "application/pdf")},
    )

    assert response.status_code == 422
    assert response.headers["content-type"].startswith("application/json")
    assert response.json()["error"]["code"] == "validation_error"


async def _async_value(value):
    return value


def test_unexpected_extraction_error_is_structured_json(
    client: TestClient,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.api.v1 import cv as cv_api

    async def fail_extraction(*_args, **_kwargs):
        raise RuntimeError("private parser detail")

    monkeypatch.setattr(cv_api, "extract_cv_sections", fail_extraction)
    response = client.post(
        "/v1/cv/extract?mode=local",
        files={"file": ("synthetic.pdf", b"synthetic", "application/pdf")},
    )

    assert response.status_code == 500
    assert response.headers["content-type"].startswith("application/json")
    assert response.json()["error"]["code"] == "internal_error"
    assert "private parser detail" not in response.text


def test_extract_ai_mode_passes_selected_model(
    client: TestClient,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.pipeline import extract as extract_pipeline
    from app.schemas.cv import CvSections

    seen: dict[str, str | None] = {}

    async def fake_ai_extract(
        _normalized_text: str,
        *,
        target_language: str = "auto",
        llm_model: str | None = None,
    ) -> CvSections:
        seen["llm_model"] = llm_model
        return CvSections(experience=["Parsed with selected model"])

    monkeypatch.setattr(extract_pipeline, "run_ai_extract_sections", fake_ai_extract)

    r = client.post(
        "/v1/cv/extract?mode=ai&llm_model=qwen%2Fqwen3.8-27b",
        files={
            "file": (
                "model-cv.docx",
                _docx_bytes(["Unstructured CV text"]),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert r.status_code == 200
    assert seen["llm_model"] == "qwen/qwen3.8-27b"
    assert r.json()["meta"]["llm_model"] == "qwen/qwen3.8-27b"


def test_hybrid_uses_ai_for_multi_column_document(
    client: TestClient,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.pipeline import extract as extract_pipeline
    from app.pipeline.preprocess import DocumentExtractResult
    from app.schemas.cv import CvSections

    async def fake_ai_extract(
        _normalized_text: str,
        *,
        target_language: str = "auto",
        llm_model: str | None = None,
    ) -> CvSections:
        return CvSections(
            experience=["05/2025 - 07/2025 | Team Member | Full project detail"],
            education=["FPT University"],
            skills=["Python"],
        )

    monkeypatch.setattr(
        extract_pipeline,
        "extract_document_text",
        lambda _filename, _data: DocumentExtractResult(
            normalized_text=(
                "Experience\nTeam Member\nEducation\nFPT University\n"
                "Skills\nPython\nProjects\nProject detail"
            ),
            engine="pymupdf",
            text_block_count=8,
            max_columns=2,
        ),
    )
    monkeypatch.setattr(extract_pipeline, "run_ai_extract_sections", fake_ai_extract)

    response = client.post(
        "/v1/cv/extract?mode=hybrid",
        files={"file": ("two-column.pdf", b"pdf", "application/pdf")},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["meta"]["parser_source"] == "ai"
    assert payload["meta"]["document"]["max_columns"] == 2
    assert "Full project detail" in payload["sections"]["experience"][0]


def test_extract_and_analyze_rejects_wrong_vertical(client: TestClient) -> None:
    response = client.post(
        "/v1/cv/extract-and-analyze?mode=local&industry_slug=law&vertical_slug=economic_law&strict_industry_match=true",
        files={
            "file": (
                "ai-cv.docx",
                _docx_bytes(
                    [
                        "Education",
                        "Computer Science student",
                        "Skills",
                        "Python, Machine Learning, Computer Vision",
                        "Projects",
                        "Built an image classification model with PyTorch",
                    ]
                ),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "validation_error"


def test_extract_and_analyze_accepts_matching_vertical(client: TestClient) -> None:
    response = client.post(
        "/v1/cv/extract-and-analyze?mode=local&industry_slug=ai_software&vertical_slug=machine_learning&target_language=en",
        files={
            "file": (
                "ai-cv.docx",
                _docx_bytes(
                    [
                        "Education",
                        "Computer Science student",
                        "Skills",
                        "Python, Machine Learning, Computer Vision",
                        "Projects",
                        "Built an image classification model with PyTorch",
                    ]
                ),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200


def test_extract_rejects_unsupported_file_type(client: TestClient) -> None:
    r = client.post(
        "/v1/cv/extract",
        files={"file": ("cv.txt", b"Experience\nBuilt a project", "text/plain")},
    )

    assert r.status_code == 422
    assert r.json()["error"]["code"] == "validation_error"
